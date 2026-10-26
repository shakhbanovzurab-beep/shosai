"""A document's text read in a bounded child process: each format, its caps, a bomb, a broken
file and a child that runs over its time."""

import io
import zipfile
from datetime import timedelta

import pytest
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

from shosai.core.documents.extraction import DocumentReader, ExtractionError
from shosai.core.documents.formats import DocumentFormat, FileKind, sniff


def pdf_bytes(text: str) -> bytes:
    """A one-page PDF with the text, its cross-reference table counted byte by byte."""
    stream = f"BT /F1 24 Tf 72 700 Td ({text}) Tj ET".encode()
    objects = [
        b"<</Type/Catalog/Pages 2 0 R>>",
        b"<</Type/Pages/Kids[3 0 R]/Count 1>>",
        b"<</Type/Page/Parent 2 0 R/MediaBox[0 0 612 792]/Contents 4 0 R"
        b"/Resources<</Font<</F1 5 0 R>>>>>>",
        b"<</Length %d>>stream\n" % len(stream) + stream + b"\nendstream",
        b"<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>",
    ]
    out, offsets = bytearray(b"%PDF-1.4\n"), []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj" % number + body + b"endobj\n"
    table = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % offset for offset in offsets)
    out += b"trailer<</Size %d/Root 1 0 R>>\nstartxref\n%d\n%%%%EOF" % (len(objects) + 1, table)
    return bytes(out)


PDF = pdf_bytes("Budget for Q3")


def docx_bytes() -> bytes:
    document = Document()
    document.add_paragraph("Northwind signs at €18 per seat.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Seats", "120"
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def xlsx_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Plan"
    sheet.append(["Month", "Revenue"])
    sheet.append(["March", 1200])
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


def pptx_bytes() -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.text_frame.text = "Launch in November"
    out = io.BytesIO()
    presentation.save(out)
    return out.getvalue()


@pytest.mark.parametrize(
    ("document", "data", "expected"),
    [
        (DocumentFormat.PDF, PDF, "Budget for Q3"),
        (DocumentFormat.DOCX, docx_bytes(), "€18 per seat"),
        (DocumentFormat.XLSX, xlsx_bytes(), "March\t1200"),
        (DocumentFormat.PPTX, pptx_bytes(), "Launch in November"),
        (DocumentFormat.RTF, rb"{\rtf1\ansi Signed {\b today}.}", "Signed today."),
        (DocumentFormat.TEXT, "Привет, Leo".encode(), "Привет, Leo"),
    ],
)
async def test_each_format_yields_its_text(
    document: DocumentFormat, data: bytes, expected: str
) -> None:
    read = await DocumentReader().text(document, data, 10_000)

    assert expected in read.text
    assert read.truncated is False


async def test_the_text_stops_at_its_limit() -> None:
    read = await DocumentReader().text(DocumentFormat.TEXT, b"a" * 5_000, 100)

    assert (len(read.text), read.truncated) == (100, True)


async def test_a_bomb_and_a_broken_file_fail_alone_with_their_code() -> None:
    bomb = io.BytesIO()
    with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"\0" * (20 * 1024**2))
    reader = DocumentReader()

    with pytest.raises(ExtractionError) as swollen:
        await reader.text(DocumentFormat.DOCX, bomb.getvalue(), 10_000)
    with pytest.raises(ExtractionError) as broken:
        await reader.text(DocumentFormat.PDF, b"%PDF-1.4\nnothing here", 10_000)
    good = await reader.text(DocumentFormat.TEXT, b"still fine", 100)

    assert (swollen.value.code, broken.value.code, good.text) == (
        "too_large",
        "malformed",
        "still fine",
    )


async def test_a_child_that_runs_over_its_time_is_killed() -> None:
    reader = DocumentReader(timeout=timedelta(milliseconds=1))

    with pytest.raises(ExtractionError) as late:
        await reader.text(DocumentFormat.PDF, PDF, 10_000)

    assert late.value.code == "timeout"


def test_the_bytes_decide_the_kind_and_the_name_only_tells_office_apart() -> None:
    zipped = docx_bytes()

    assert sniff("photo.txt", b"\x89PNG\r\n\x1a\n...").media_type == "image/png"
    assert sniff("scan.png", PDF).document is DocumentFormat.PDF
    assert sniff("report.docx", zipped).document is DocumentFormat.DOCX
    assert sniff("report.zip", zipped).kind is FileKind.REJECTED
    assert sniff("notes.md", b"# Plan").document is DocumentFormat.TEXT
    assert sniff("notes.md", b"\xff\xfe\x00bad").kind is FileKind.REJECTED
    assert sniff("program.exe", b"MZ\x90\x00").kind is FileKind.REJECTED
    assert sniff("empty.txt", b"").kind is FileKind.REJECTED
